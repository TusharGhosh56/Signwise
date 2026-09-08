import io
import re
from typing import Dict, Any, List
from docx import Document


class DOCXParseError(Exception):
    pass


def extract_text_from_docx_bytes(docx_bytes: bytes) -> Dict[str, Any]:
    """
    Extracts structured plain text and metadata from DOCX byte stream.
    Reads paragraphs, headings, and table cells.
    Returns:
        {
            "full_text": str,
            "paragraph_count": int,
            "table_count": int,
            "word_count": int,
            "title": str | None
        }
    """
    try:
        stream = io.BytesIO(docx_bytes)
        doc = Document(stream)

        text_chunks: List[str] = []

        # Read core properties if available
        meta_title = None
        if doc.core_properties and doc.core_properties.title:
            meta_title = str(doc.core_properties.title).strip()

        # Extract all paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # Clean control characters
                cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()
                if cleaned:
                    text_chunks.append(cleaned)

        # Extract tables
        table_count = len(doc.tables)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_chunks.append(" | ".join(row_cells))

        full_text = "\n\n".join(text_chunks)
        word_count = len(re.findall(r"\w+", full_text))

        if word_count < 10:
            raise DOCXParseError(
                "Could not extract sufficient text from this Word document. "
                "The file appears to be empty or contains unsupported embedded objects."
            )

        return {
            "full_text": full_text,
            "paragraph_count": len(doc.paragraphs),
            "table_count": table_count,
            "word_count": word_count,
            "title": meta_title,
        }

    except DOCXParseError:
        raise
    except Exception as e:
        raise DOCXParseError(f"Failed to read Word document: {str(e)}")
